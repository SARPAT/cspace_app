"""DocQuery RAG engine: ingestion, hybrid retrieval, prompting and generation.

Pipeline
    load -> chunk -> index (ChromaDB dense vectors + BM25 sparse index)
    question -> hybrid top-k retrieval -> grounded prompt -> LLM answer with [n] citations
If no LLM key is configured, an extractive generator answers from the retrieved text.
"""
from __future__ import annotations

import io
import json
import math
import re
import sys
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path

# Streamlit Community Cloud ships an old sqlite3; ChromaDB needs >= 3.35.
try:  # pragma: no cover - only relevant on Linux hosts with pysqlite3-binary installed
    __import__("pysqlite3")
    sys.modules["sqlite3"] = sys.modules.pop("pysqlite3")
except ImportError:
    pass

from rank_bm25 import BM25Okapi

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

# --------------------------------------------------------------------------- text utils
STOP = set(
    "a an the of to in and or for on by with at from as is are was were be been being it its this that these "
    "those which who whom what when where how why do does did can could shall should will would may might must "
    "not no any all every each per under into than then there their they them his her he she we our you your i "
    "me my also such other only about over after before up out if so but have has had give gives given get gets "
    "much many some one two".split()
)


def stem(w: str) -> str:
    if len(w) > 5 and w.endswith("ies"):
        return w[:-3] + "y"
    if len(w) > 5 and w.endswith("ing"):
        return w[:-3]
    if len(w) > 4 and w.endswith("ed"):
        return w[:-2]
    if len(w) > 4 and w.endswith("es") and not w.endswith("ses"):
        return w[:-2]
    if len(w) > 3 and w.endswith("s") and not w.endswith("ss"):
        return w[:-1]
    return w


def tokens(text: str) -> list[str]:
    words = re.findall(r"[a-z0-9]+", text.lower().replace("'", "").replace("’", ""))
    return [stem(w) for w in words if w not in STOP]


SENT_SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z0-9\"'(\[])")


def split_sentences(text: str) -> list[str]:
    return [s.strip() for s in SENT_SPLIT.split(text) if s.strip()]


# --------------------------------------------------------------------------- data model
@dataclass
class Document:
    doc_id: str
    title: str
    category: str
    doc_type: str
    chunks: list[str]


@dataclass
class Chunk:
    chunk_id: str
    doc: Document
    idx: int
    text: str
    words: int = field(init=False)

    def __post_init__(self):
        self.words = len(self.text.split())


@dataclass
class Hit:
    chunk: Chunk
    score: float
    bm25: float = 0.0
    dense: float = 0.0
    header: bool = False


# --------------------------------------------------------------------------- ingestion
def chunk_text(text: str, chunk_size: int = 700, overlap: int = 100) -> list[str]:
    """Paragraph/sentence-aware splitting (LangChain RecursiveCharacterTextSplitter)."""
    text = re.sub(r"[ \t]+", " ", text.replace("\r", ""))
    try:
        from langchain_text_splitters import RecursiveCharacterTextSplitter

        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size, chunk_overlap=overlap, separators=["\n\n", "\n", ". ", " ", ""]
        )
        parts = splitter.split_text(text)
    except ImportError:  # simple fallback
        parts, buf = [], ""
        for para in re.split(r"\n\s*\n", text):
            if len(buf) + len(para) > chunk_size and buf:
                parts.append(buf)
                buf = ""
            buf = (buf + " " + para).strip()
        if buf:
            parts.append(buf)
    cleaned = [re.sub(r"\s+", " ", p).strip() for p in parts]
    return [c for c in cleaned if len(c.split()) >= 5]


def extract_text(name: str, data: bytes) -> tuple[str, str]:
    """Return (text, doc_type) for an uploaded file."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(p.extract_text() or "") for p in reader.pages[:150]]
        return "\n\n".join(pages), "PDF"
    if ext == "docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        return "\n\n".join(p.text for p in d.paragraphs if p.text.strip()), "Word"
    text = data.decode("utf-8", errors="ignore")
    if ext in ("html", "htm"):
        text = re.sub(r"(?is)<(script|style).*?</\1>", " ", text)
        text = re.sub(r"<[^>]+>", " ", text)
        return text, "HTML"
    return text, "Markdown" if ext in ("md", "markdown") else "TXT"


def load_builtin() -> tuple[list[Document], list[dict]]:
    raw = json.loads((DATA_DIR / "corpus.json").read_text(encoding="utf-8"))
    docs = [
        Document(d["id"], d["title"], d["cat"], d["type"], d["chunks"]) for d in raw["documents"]
    ]
    return docs, raw["benchmark"]


# --------------------------------------------------------------------------- index
class HybridIndex:
    """BM25 (sparse) + ChromaDB (dense, all-MiniLM-L6-v2) hybrid retriever."""

    def __init__(self, use_dense: bool = True):
        self.chunks: list[Chunk] = []
        self.docs: list[Document] = []
        self.bm25: BM25Okapi | None = None
        self.collection = None
        self.dense_error: str | None = None
        if use_dense:
            try:
                import chromadb

                client = chromadb.EphemeralClient()
                self.collection = client.create_collection(
                    name=f"docquery_{uuid.uuid4().hex[:8]}", metadata={"hnsw:space": "cosine"}
                )
            except Exception as e:  # noqa: BLE001 - keep the app usable without dense search
                self.dense_error = str(e)

    @property
    def dense_on(self) -> bool:
        return self.collection is not None

    def add(self, docs: list[Document]) -> None:
        new = []
        for d in docs:
            self.docs.append(d)
            for i, text in enumerate(d.chunks):
                c = Chunk(f"CHK{len(self.chunks) + 1:05d}", d, i, text)
                self.chunks.append(c)
                new.append(c)
        self.bm25 = BM25Okapi([tokens(c.text) or ["_"] for c in self.chunks], k1=1.4, b=0.75)
        if self.collection is not None and new:
            try:
                for s in range(0, len(new), 64):
                    batch = new[s : s + 64]
                    self.collection.add(
                        ids=[c.chunk_id for c in batch],
                        documents=[c.text for c in batch],
                        metadatas=[{"doc_id": c.doc.doc_id} for c in batch],
                    )
            except Exception as e:  # noqa: BLE001 - e.g. embedding model download blocked
                self.dense_error = str(e)
                self.collection = None

    def retrieve(self, question: str, k: int = 4) -> list[Hit]:
        if not self.chunks:
            return []
        q = tokens(question)
        bm = self.bm25.get_scores(q) if q else [0.0] * len(self.chunks)
        top_bm = max(bm) if len(bm) else 0
        dense = [0.0] * len(self.chunks)
        if self.collection is not None:
            try:
                n = min(len(self.chunks), max(k * 4, 12))
                res = self.collection.query(query_texts=[question], n_results=n)
                pos = {c.chunk_id: i for i, c in enumerate(self.chunks)}
                for cid, dist in zip(res["ids"][0], res["distances"][0]):
                    dense[pos[cid]] = max(0.0, 1.0 - float(dist))
            except Exception as e:  # noqa: BLE001
                self.dense_error = str(e)
        hits = []
        for i, c in enumerate(self.chunks):
            b = (bm[i] / top_bm) if top_bm > 0 else 0.0
            if b <= 0 and dense[i] < 0.35:
                continue
            score = 0.6 * b + 0.4 * dense[i] if self.dense_on else b
            hits.append(Hit(c, score, bm25=float(bm[i]), dense=dense[i]))
        hits.sort(key=lambda h: h.score, reverse=True)
        return hits[:k]

    def with_headers(self, hits: list[Hit]) -> list[Hit]:
        """Also pass the opening chunk(s) of the top documents: title, authors, abstract."""
        out, have = list(hits), {h.chunk.chunk_id for h in hits}
        seen_docs = []
        for h in hits[:2]:
            if h.chunk.doc not in seen_docs:
                seen_docs.append(h.chunk.doc)
        for d in seen_docs:
            n = 2 if d.category == "Your uploads" else 1
            for c in self.chunks:
                if c.doc is d and c.idx < n and c.chunk_id not in have:
                    have.add(c.chunk_id)
                    out.append(Hit(c, 0.0, header=True))
        return out


# --------------------------------------------------------------------------- prompting
def build_prompt(question: str, hits: list[Hit]) -> str:
    ctx = "\n\n".join(
        f"[{i}] ({h.chunk.doc.title}, {h.chunk.chunk_id}"
        f"{', opening of the document: usually title, authors, abstract' if h.header else ''})\n{h.chunk.text}"
        for i, h in enumerate(hits, 1)
    )
    return f"""You are a document question-answering assistant in a RAG system. Answer the QUESTION using ONLY the numbered CONTEXT chunks below. Rules:
- Do not use outside knowledge. If the context does not contain the answer, reply exactly: "I could not find this in the documents."
- Cite the chunk numbers you used in square brackets after each claim, e.g. [1] or [2][3].
- Start with a direct one-sentence answer, then add supporting detail only if it helps.
- When the answer is a set of items (for example authors, steps, features, conditions), list EVERY item from the context as a bullet line starting with "- ", one item per line.
- Use **bold** for the key names, numbers or terms. No headings, no preamble. Keep it under 180 words.

CONTEXT:
{ctx}

QUESTION: {question}

ANSWER:"""


# --------------------------------------------------------------------------- LLM providers (OpenAI-compatible APIs)
PROVIDERS = {
    "Groq (free tier)": {"base_url": "https://api.groq.com/openai/v1", "model": "llama-3.3-70b-versatile", "secret": "GROQ_API_KEY"},
    "Grok (xAI)": {"base_url": "https://api.x.ai/v1", "model": "grok-3-mini", "secret": "XAI_API_KEY"},
    "OpenAI": {"base_url": None, "model": "gpt-4o-mini", "secret": "OPENAI_API_KEY"},
    "Ollama (local)": {"base_url": "http://localhost:11434/v1", "model": "llama3.2", "secret": None},
}


def stream_answer(provider: str, api_key: str, model: str, prompt: str):
    """Yield answer text pieces from an OpenAI-compatible chat endpoint."""
    from openai import OpenAI

    cfg = PROVIDERS[provider]
    client = OpenAI(api_key=api_key or "ollama", base_url=cfg["base_url"], timeout=60)
    stream = client.chat.completions.create(
        model=model,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.1,
        max_tokens=500,
        stream=True,
    )
    for part in stream:
        if part.choices and part.choices[0].delta and part.choices[0].delta.content:
            yield part.choices[0].delta.content


# --------------------------------------------------------------------------- extractive fallback
AFFIL = re.compile(
    r"(Department|Dept\.|School of|Faculty of|University|Institute|College|Laborator|Centre|Center for|"
    r"[\w.+-]+@[\w-]+\.|Abstract|ABSTRACT|E-?mail)"
)
NOT_NAME = set(
    "india engineering science sciences computer technology research department university college institute "
    "school abstract email paper journal conference international national proceedings ieee acm vol volume "
    "analysis text texts image images detection learning retrieval network networks model models system systems "
    "approach study data deep legal based using towards via framework survey review method methods multi modal "
    "neural language large generation augmented phishing security secure".split()
)
NAME_WORD = re.compile(r"^[A-Z][a-z'’-]+$")
INITIAL = re.compile(r"^[A-Z]\.$")


def find_authors(index: HybridIndex, hits: list[Hit]) -> list[tuple[str, int]]:
    if not hits:
        return []
    target = hits[0].chunk.doc
    doc_text = " ".join(c.text for c in index.chunks if c.doc is target)
    lower = set(re.findall(r"\b[a-z][a-z'’-]+\b", doc_text))
    found, seen = [], set()
    for hi, h in enumerate(hits, 1):
        if h.chunk.idx > 1 or h.chunk.doc is not target:
            continue
        t = h.chunk.text[:1500]
        for m in AFFIL.finditer(t):
            seg = t[max(0, m.start() - 90) : m.start()]
            for piece in re.split(r",|;|&|\band\b|\d|\*|†|‡|\|", seg):
                words, run = piece.split(), []
                for w in reversed(words):
                    if len(run) >= 4:
                        break
                    ok = INITIAL.match(w) or (
                        NAME_WORD.match(w) and w.lower() not in lower and w.lower() not in NOT_NAME
                    )
                    if not ok:
                        break
                    run.insert(0, w)
                if len(run) >= 2 and not INITIAL.match(run[-1]):
                    name = " ".join(run)
                    if name not in seen:
                        seen.add(name)
                        found.append((name, hi))
    return found


def _clip(s: str) -> str:
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= 300 else s[:297].rsplit(" ", 1)[0] + "…"


def extractive_answer(index: HybridIndex, question: str, hits: list[Hit]) -> str:
    if not hits:
        return "I could not find this in the documents."
    if re.search(r"\b(authors?|written by|who wrote|wrote)\b", question, re.I):
        authors = find_authors(index, hits)
        if authors:
            title = hits[0].chunk.doc.title
            head = f"**{title}** lists " + ("one author:" if len(authors) == 1 else f"{len(authors)} authors:")
            return head + "\n" + "\n".join(f"- **{n}** [{i}]" for n, i in authors)
    qs = set(tokens(question))
    idf = {}
    n = len(index.chunks)
    for t in qs:
        df = sum(1 for c in index.chunks if t in tokens(c.text))
        idf[t] = math.log(1 + (n - df + 0.5) / (df + 0.5))
    sents = []
    for hi, h in enumerate(hits, 1):
        for si, s in enumerate(split_sentences(h.chunk.text)):
            st = tokens(s)
            if len(st) < 3:
                continue
            sc = sum(idf.get(t, 0) for t in set(st) if t in qs)
            sc = sc * (1 + h.score) / (1 + math.log(1 + len(st)) * 0.15)
            if sc > 0:
                sents.append((sc, hi, si, _clip(s)))
    if not sents:
        return "I could not find this in the documents."
    sents.sort(reverse=True)
    best, keep = sents[0][0], []
    for x in sents:
        if len(keep) >= 4 or x[0] < best * 0.3:
            break
        if all(k[3] != x[3] for k in keep):
            keep.append(x)
    lead, rest = keep[0], sorted(keep[1:], key=lambda x: (x[1], x[2]))
    out = f"**Best match:** {lead[3]} [{lead[1]}]"
    if rest:
        out += "\n\n**Related passages:**\n" + "\n".join(f"- {x[3]} [{x[1]}]" for x in rest)
    return out


# --------------------------------------------------------------------------- benchmark
def run_benchmark(index: HybridIndex, benchmark: list[dict], k: int = 5) -> dict:
    rows, h1, h3, rr, lat = [], 0, 0, 0.0, []
    for item in benchmark:
        t = time.perf_counter()
        hits = index.retrieve(item["question"], k)
        lat.append((time.perf_counter() - t) * 1000)
        ids = [h.chunk.doc.doc_id for h in hits]
        pos = ids.index(item["expected_doc"]) if item["expected_doc"] in ids else -1
        h1 += pos == 0
        h3 += 0 <= pos < 3
        rr += 1 / (pos + 1) if pos >= 0 else 0
        rows.append(
            {
                "Question": item["question"],
                "Expected doc": item["expected_doc"],
                "Top-1 doc": ids[0] if ids else "none",
                "Rank of expected": f"#{pos + 1}" if pos >= 0 else "not in top 5",
                "Top score": round(hits[0].score, 3) if hits else 0.0,
            }
        )
    n = max(1, len(benchmark))
    return {"hit1": h1 / n, "hit3": h3 / n, "mrr": rr / n, "latency_ms": sum(lat) / n, "rows": rows}
